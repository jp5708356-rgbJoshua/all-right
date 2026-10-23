"""合同规则检查（纯正则，不做语义分析、不调模型、不给修改建议）。

规则在 rules/contract_rules.yaml 里维护，本模块只负责：
    读文本 -> 逐条正则匹配 -> 输出带原文位置的风险清单

所有命中结果一律 needs_human_review = True，代码里硬置，没有关闭入口。
"""

from __future__ import annotations

import io
import re
from pathlib import Path

import yaml

ROOT = Path(__file__).resolve().parent.parent
RULES_PATH = ROOT / "rules" / "contract_rules.yaml"

REQUIRED_KEYS = ("id", "category", "level", "pattern", "message")
LEVELS = ("高", "中", "低")
SUPPORTED_SUFFIXES = (".txt", ".md", ".docx", ".pdf")

# exclude 正则匹配命中位置前后多少字，判定是否属于合理表述
EXCLUDE_WINDOW = 30


class RuleError(ValueError):
    """规则文件本身有问题。"""


def load_rules(path: str | Path | None = None) -> list[dict]:
    rules_path = Path(path) if path else RULES_PATH
    if not rules_path.exists():
        raise FileNotFoundError(f"规则文件不存在：{rules_path}")

    data = yaml.safe_load(rules_path.read_text(encoding="utf-8")) or {}
    rules = data.get("rules") or []
    if not rules:
        raise RuleError("规则文件里没有 rules 列表")

    for idx, rule in enumerate(rules, 1):
        missing = [key for key in REQUIRED_KEYS if not rule.get(key)]
        if missing:
            raise RuleError(f"第 {idx} 条规则缺少字段：{'、'.join(missing)}")
        if rule["level"] not in LEVELS:
            raise RuleError(f"规则 {rule['id']} 的 level 只能是 高 / 中 / 低")
        try:
            re.compile(rule["pattern"])
        except re.error as exc:
            raise RuleError(f"规则 {rule['id']} 的正则写错了：{exc}") from exc

        if rule.get("exclude"):
            try:
                re.compile(rule["exclude"])
            except re.error as exc:
                raise RuleError(f"规则 {rule['id']} 的 exclude 正则写错了：{exc}") from exc

    return rules


def extract_text(source, filename: str | None = None) -> str:
    """从路径 / bytes / 上传文件对象里取纯文本。"""
    if isinstance(source, (str, Path)):
        path = Path(source)
        name = path.name
        suffix = path.suffix.lower()
        data = path.read_bytes()
    else:
        name = filename or getattr(source, "name", "") or ""
        suffix = Path(name).suffix.lower()
        if isinstance(source, bytes):
            data = source
        else:
            data = source.read()
            if isinstance(data, str):
                return data

    if suffix in (".txt", ".md", ""):
        return data.decode("utf-8", errors="replace")

    if suffix == ".docx":
        try:
            from docx import Document
        except ImportError as exc:
            raise ValueError("读 .docx 需要 python-docx，请先 pip install python-docx") from exc
        document = Document(io.BytesIO(data))
        parts = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.append("\t".join(cell.text for cell in row.cells))
        return "\n".join(parts)

    if suffix == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise ValueError("读 .pdf 需要 pypdf，请先 pip install pypdf") from exc
        reader = PdfReader(io.BytesIO(data))
        return "\n".join((page.extract_text() or "") for page in reader.pages)

    raise ValueError(f"不支持的文件类型 {suffix or '（无扩展名）'}，请用 .txt / .docx / .pdf")


def _snippet(text: str, start: int, end: int, width: int = 40) -> str:
    left = max(0, start - width)
    right = min(len(text), end + width)
    piece = text[left:right].replace("\n", " ").replace("\r", " ")
    prefix = "…" if left > 0 else ""
    suffix = "…" if right < len(text) else ""
    return f"{prefix}{piece}{suffix}"


def check(text: str, rules: list[dict] | None = None) -> dict:
    """对纯文本跑一遍规则。

    规则可选带 `exclude` 正则：命中位置的上下文里如果出现 exclude 内容，
    视为合理表述跳过。例如「甲方自身原因造成的损失，乙方不承担任何责任」
    是合理免责，不该和「乙方概不负责」报成同一件事。
    """
    if not text or not text.strip():
        raise ValueError("合同内容为空")

    active_rules = rules if rules is not None else load_rules()
    results: list[dict] = []
    suppressed = 0

    for rule in active_rules:
        pattern = re.compile(rule["pattern"], re.MULTILINE)
        exclude = re.compile(rule["exclude"], re.MULTILINE) if rule.get("exclude") else None

        for match in pattern.finditer(text):
            if exclude is not None:
                left = max(0, match.start() - EXCLUDE_WINDOW)
                right = min(len(text), match.end() + EXCLUDE_WINDOW)
                if exclude.search(text[left:right]):
                    suppressed += 1
                    continue

            results.append(
                {
                    "rule_id": rule["id"],
                    "category": rule["category"],
                    "level": rule["level"],
                    "message": rule["message"],
                    "matched_text": match.group(0),
                    "char_offset": match.start(),
                    "line_no": text.count("\n", 0, match.start()) + 1,
                    "snippet": _snippet(text, match.start(), match.end()),
                    # 硬置：任何命中都必须人工复核，不给关闭入口
                    "needs_human_review": True,
                }
            )

    order = {level: idx for idx, level in enumerate(LEVELS)}
    results.sort(key=lambda r: (order.get(r["level"], 9), r["char_offset"]))

    return {
        "text_length": len(text),
        "rule_count": len(active_rules),
        "hit_count": len(results),
        "by_level": {level: sum(1 for r in results if r["level"] == level) for level in LEVELS},
        "results": results,
        "no_hit": not results,
        "suppressed": suppressed,
    }


def check_file(source, filename: str | None = None, rules: list[dict] | None = None) -> dict:
    text = extract_text(source, filename)
    report = check(text, rules)
    report["filename"] = filename or (Path(source).name if isinstance(source, (str, Path)) else "")
    return report
